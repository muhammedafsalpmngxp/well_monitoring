# Ported from twise-agentic db_assistant/src/tools/export_formats.py (team repo); see app/askdata/README.md.
"""Turn (columns, rows) into a downloadable file. No LLM, no database — pure formatting.

Two formats, both spreadsheet/document native:
  * xlsx  — the data, every row (subject only to the .xlsx format's own sheet limit)
  * docx  — the answer's narrative, then the data as a real Word table

Both are given the SAME rows, so an Excel and a Word export of one answer can never disagree.

Heavy imports (openpyxl / python-docx) are done INSIDE each function rather than at module scope,
so a missing one degrades to a clear message for that single format instead of breaking the whole
API at import time.
"""
from __future__ import annotations

import datetime as dt
import io
import re
from typing import Any

# One place decides how a value is written, so both formats render a date/None identically.
_NULL = ""


def cell(value: Any) -> str:
    if value is None:
        return _NULL
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def header(name: str) -> str:
    """Column name in business language: created_at -> Created At."""
    return str(name).replace("_", " ").strip().title()


# A capped export must SAY it was capped — the same honesty rule the chat answer follows.
def cap_note(capped: bool, n_rows: int, limit: int) -> str:
    if not capped:
        return ""
    return (
        f"NOTE: this export was limited to {limit:,} rows; the full result contains more than the "
        f"{n_rows:,} rows shown here."
    )


def unverified_note(verified: bool) -> str:
    if verified:
        return ""
    return (
        "NOTE: the figures in this file could not be fully verified and should be treated as "
        "indicative."
    )


# ── Answer narrative -> document blocks (Word only) ──────────────────────────────
# The Word export puts the answer above the data table. The answer, however, OFTEN ALREADY
# CONTAINS that same data as a markdown table (the Synthesizer is told to format results as a
# table), which produced the whole result twice in one document: once as unreadable pipe-
# separated text, then again as the real table.
#
# So the narrative is filtered: anything that is tabular markup is dropped, because the real
# table renders it properly two paragraphs later. Prose, headings and bullets are kept — they are
# the part a table cannot express, and the part that makes the document worth reading.

# A markdown separator row: "--- | ---: | ---" or "|:--|--:|".
_MD_SEPARATOR = re.compile(r"^[\s|:-]*-{2,}[\s|:-]*$")
# Inline emphasis/code markers, which would otherwise show as literal ** and ` in Word.
_MD_BOLD = re.compile(r"\*\*(.+?)\*\*|__(.+?)__")
_MD_ITALIC = re.compile(r"(?<!\*)\*(?!\s)([^*]+?)(?<!\s)\*(?!\*)")
_MD_CODE = re.compile(r"`([^`]+)`")
_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_MD_BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")
_MD_NUMBER = re.compile(r"^\s*\d+[.)]\s+(.*)$")
# The chat answer may point at the Export control ("...download using the Export option below").
# Inside the exported file that sentence is nonsense: this file IS the download, and it has no
# Export button. Dropped rather than shipped as a dangling instruction.
_EXPORT_HINT = re.compile(r"(export option|can be downloaded|download(ed)? (it |them |the )?using)", re.I)


def _tabular_lines(lines: list[str]) -> set[int]:
    """Indexes of the lines that are table markup rather than prose.

    Decided over the WHOLE answer, not line by line, because the pipe count alone is ambiguous.
    A per-line rule of ">= 2 pipes" caught wide tables but silently missed TWO-COLUMN ones -
    "Al Tasnim Enterprises LLC | 5193" carries a single pipe - so an entire company/count table
    was reproduced as prose above the real table.

    The structural signal is repetition: one line containing a pipe may well be a sentence, but
    several of them are a table. So pipe-bearing lines are only treated as tabular when at least
    two of them exist, which needs no knowledge of the column count and no hardcoded headings.
    """
    pipe_rows = [i for i, l in enumerate(lines) if l.strip() and "|" in l]
    separators = {
        i for i, l in enumerate(lines)
        if l.strip() and _MD_SEPARATOR.match(l.strip()) and "-" in l
    }
    # A lone pipe line is prose (e.g. "Status: Active | Inactive"); two or more form a table.
    if len(pipe_rows) < 2:
        return separators | {i for i in pipe_rows if lines[i].count("|") >= 2}
    return separators | set(pipe_rows)


def _inline(text: str) -> str:
    text = _MD_BOLD.sub(lambda m: m.group(1) or m.group(2) or "", text)
    text = _MD_ITALIC.sub(r"\1", text)
    text = _MD_CODE.sub(r"\1", text)
    return text.strip()


def narrative_blocks(answer: str) -> list[tuple[str, str]]:
    """(kind, text) for each line worth putting above the table.

    kind is "heading" | "bullet" | "number" | "body". Table markup and the export pointer are
    dropped. Exposed (not private) so the behaviour is directly testable.
    """
    blocks: list[tuple[str, str]] = []
    lines = [l.rstrip() for l in (answer or "").splitlines()]
    tabular = _tabular_lines(lines)
    for idx, line in enumerate(lines):
        if not line.strip():
            continue
        if idx in tabular:
            continue                      # the real table covers this
        if _EXPORT_HINT.search(line):
            continue                      # meaningless inside the downloaded file itself
        m = _MD_HEADING.match(line)
        if m:
            text = _inline(m.group(2))
            if text:
                blocks.append(("heading", text))
            continue
        m = _MD_BULLET.match(line)
        if m:
            text = _inline(m.group(1))
            if text:
                blocks.append(("bullet", text))
            continue
        m = _MD_NUMBER.match(line)
        if m:
            text = _inline(m.group(1))
            if text:
                blocks.append(("number", text))
            continue
        text = _inline(line)
        if text:
            blocks.append(("body", text))
    return blocks


# ── Column sizing ───────────────────────────────────────────────────────────────
# How many rows to measure when deciding column widths. A full pass would defeat the point of
# streaming a million-row export, and the first few hundred rows are representative enough to
# stop a column being visibly too narrow.
WIDTH_SAMPLE_ROWS = 200
MIN_WIDTH_CHARS = 8
MAX_WIDTH_CHARS = 46


def content_widths(columns: list[str], rows: list[list]) -> list[int]:
    """Width in characters per column, wide enough for the HEADER and a sample of the values.

    Sizing on the header alone was not enough: "Eng Completion Date" is 19 characters, so a
    header-only guess still clipped it, and date columns looked cramped next to it.
    """
    widths = [len(header(c)) for c in columns]
    for row in rows[:WIDTH_SAMPLE_ROWS]:
        for i, value in enumerate(row):
            if i < len(widths):
                widths[i] = max(widths[i], len(cell(value)))
    return [max(MIN_WIDTH_CHARS, min(MAX_WIDTH_CHARS, w + 2)) for w in widths]


# Excel rejects these characters in a SHEET NAME and openpyxl raises ValueError on them.
# The set is taken from openpyxl's own INVALID_TITLE_REGEX, not guessed: backslash * ? : / [ ]
#
# This mattered in practice: the sheet name is derived from the user's question, and a question
# almost always ends in "?" - so nearly every real Excel export crashed with
# "Invalid character ? found in sheet title" and the endpoint returned HTTP 500.
#
# A plain set, not a regex character class: getting a backslash into a class needs double
# escaping, which is easy to get subtly wrong (the first version of this fix silently omitted
# the backslash and still looked correct).
_INVALID_SHEET_CHARS = frozenset('\\*?:/[]')


def sheet_title(title: str) -> str:
    """A sheet name Excel will accept: forbidden characters out, <=31 chars, never empty.

    Forbidden characters become a space rather than vanishing, so "wells/projects" reads as
    "wells projects" instead of "wellsprojects"; the split/join then collapses the double gaps.
    Excel also refuses an empty name, so a title made only of forbidden characters falls back
    to "Export" rather than raising.
    """
    cleaned = "".join(" " if c in _INVALID_SHEET_CHARS else c for c in (title or ""))
    return " ".join(cleaned.split())[:31].strip() or "Export"


# ── Excel ───────────────────────────────────────────────────────────────────────

# Hard limit of the .xlsx format itself, not a policy choice.
XLSX_MAX_ROWS = 1_048_576


def to_xlsx(columns: list[str], rows: list[list], notes: list[str], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension

    # write_only streams rows out rather than holding a cell object per value in memory.
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet_title(title))

    # MUST come before the first append(): in write_only mode the sheet header is written as soon
    # as rows start streaming, so column widths set afterwards are silently DISCARDED (they were —
    # every column came out at the default width and long headers were clipped).
    for i, width in enumerate(content_widths(columns, rows), 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter] = ColumnDimension(ws, index=letter, width=width)
    # Keep the header visible while scrolling a long export.
    ws.freeze_panes = "A2"

    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F3864")
    wrap = Alignment(vertical="center", wrap_text=True)
    head = []
    for name in columns:
        c = WriteOnlyCell(ws, value=header(name))
        c.font = bold
        c.fill = fill
        c.alignment = wrap
        head.append(c)
    ws.append(head)

    # Leave room for the header row and the trailing notes.
    body_limit = XLSX_MAX_ROWS - 1 - len(notes) - 2
    written = 0
    for row in rows:
        if written >= body_limit:
            notes = list(notes) + [
                f"NOTE: truncated at {body_limit:,} rows — one .xlsx sheet cannot hold more than "
                f"{XLSX_MAX_ROWS:,} rows."
            ]
            break
        ws.append([cell(v) for v in row])
        written += 1

    for note in [n for n in notes if n]:
        ws.append([])
        ws.append([note])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ── Word ────────────────────────────────────────────────────────────────────────

# python-docx builds the whole document tree in memory, and a table of tens of thousands of rows
# is both enormous and unopenable — so Word is capped well below the Excel limit and says so.
DOCX_MAX_ROWS = 5_000


def to_docx(
    columns: list[str], rows: list[list], notes: list[str], title: str, answer: str
) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Emu, Inches, Pt

    doc = Document()
    section = doc.sections[0]
    # Landscape: these result sets are wide (the well query returns 14+ columns).
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    # The default 1.25in side margins waste 2.5in of an 11in landscape page — on a 14-column
    # table that is the difference between readable and cramped. 0.5in gives 10in of usable
    # width instead of 8.5in.
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.6)

    doc.add_heading(title or "Data export", level=1)

    # The narrative only — the data itself follows as a real table, so any markdown table in the
    # answer is filtered out here rather than printed twice.
    for kind, text in narrative_blocks(answer):
        if kind == "heading":
            doc.add_heading(text, level=2)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        elif kind == "number":
            doc.add_paragraph(text, style="List Number")
        else:
            doc.add_paragraph(text)

    shown = rows[:DOCX_MAX_ROWS]
    extra = list(notes)
    if len(rows) > DOCX_MAX_ROWS:
        extra.append(
            f"NOTE: this document shows the first {DOCX_MAX_ROWS:,} of {len(rows):,} rows. "
            "Export as Excel for the complete data."
        )

    if columns:
        doc.add_paragraph("")
        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        # Explicit widths, so Word does not divide the page evenly and leave a date column as
        # roomy as a remarks column. autofit must be off or Word recomputes and ignores them.
        table.autofit = False
        # 99.5% of the printable width, not 100%: Word stores widths in twips, and the rounding
        # on the way in can land a couple of twips past the margin, which makes it reflow the
        # whole table. The 0.5% slack costs nothing visible and removes that risk entirely.
        printable = int(
            (section.page_width - section.left_margin - section.right_margin) * 0.995
        )
        shares = content_widths(columns, shown)
        total = sum(shares) or 1
        col_widths = [Emu(int(printable * s / total)) for s in shares]
        for i, width in enumerate(col_widths):
            table.columns[i].width = width

        # Font size follows column count: 7pt is right for a 14-column table but needlessly
        # small for a 3-column one.
        size = Pt(9) if len(columns) <= 6 else Pt(8) if len(columns) <= 10 else Pt(7)

        for i, name in enumerate(columns):
            c = table.rows[0].cells[i]
            c.width = col_widths[i]          # python-docx needs the width on CELLS too
            run = c.paragraphs[0].add_run(header(name))
            run.bold = True
            run.font.size = size
        # Repeat the header on every page — a 5,000-row table spans many pages, and a table whose
        # columns are unlabelled after page 1 is unusable.
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)

        for row in shown:
            cells = table.add_row().cells
            for i, value in enumerate(row):
                if i < len(cells):
                    cells[i].width = col_widths[i]
                    run = cells[i].paragraphs[0].add_run(cell(value))
                    run.font.size = size

    for note in [n for n in extra if n]:
        doc.add_paragraph("")
        # Italic must be set on a RUN: assigning .italic to the Paragraph object silently does
        # nothing (it just adds a stray Python attribute).
        p = doc.add_paragraph()
        run = p.add_run(note)
        run.italic = True
        run.font.size = Pt(8)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ── Registry ────────────────────────────────────────────────────────────────────
# Content type and extension per format, so main.py never hardcodes a MIME string.
FORMATS: dict[str, tuple[str, str]] = {
    "xlsx": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    ),
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
}
