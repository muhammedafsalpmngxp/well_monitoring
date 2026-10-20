"""The Excel workbook: a summary sheet, then one sheet per rule with findings, STREAMED from each rule's spool file
into write-only worksheets (memory stays flat however many rows a rule found). Ported from the original."""
from __future__ import annotations

import io

from app.features.anomaly.config import cfg
from app.features.anomaly.report import spool as spool_store
from app.features.anomaly.report.format import (
    SEVERITY_COLOURS,
    content_widths,
    header,
    number_or_text,
    sheet_names,
)

XLSX_MAX_ROWS = 1_048_576          # the hard limit of the .xlsx format itself, not a policy
_HEADER_FILL = "1F3864"
_TABLE_HEADINGS = ("Rule", "What was found", "Category", "Severity", "Records affected", "Of examined", "Share %")


def _write_summary_sheet(wb, state) -> None:
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension

    settings = cfg()
    ws = wb.create_sheet(title="Summary")
    for i, width in enumerate((14, 52, 24, 12, 16, 16, 12), 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter] = ColumnDimension(ws, index=letter, width=width)
    ws.freeze_panes = "A2"

    def styled(text, bold=False, colour=None):
        c = WriteOnlyCell(ws, value=text)
        if bold:
            c.font = Font(bold=True, color="FFFFFF" if colour == _HEADER_FILL else "000000")
        if colour:
            c.fill = PatternFill("solid", fgColor=colour)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        return c

    def finding_rows(rows):
        ws.append([styled(t, bold=True, colour=_HEADER_FILL) for t in _TABLE_HEADINGS])
        for row in rows:
            ws.append([row["rule_id"], row["title"], row["category"], styled(row["severity"]), row["anomaly_count"],
                       row["scope_total"], round(float(row["anomaly_pct"]), 2)])

    totals = state.get("totals") or {}
    ws.append([styled(settings.report_title, bold=True)])
    ws.append([f"{state.get('database_name') or 'database'} - run {state.get('run_id', '')}"])
    ws.append([f"Overall score: {state.get('score', 0)} / 100"])
    ws.append([state.get("score_basis", "")])
    ws.append([])
    ws.append([f"Probes run: {totals.get('probes_run', 0)}", f"With findings: {totals.get('probes_with_findings', 0)}",
               f"Clean: {totals.get('probes_clean', 0)}", f"Failed: {totals.get('probes_failed', 0)}",
               f"Records examined: {totals.get('records_examined', 0):,}",
               f"Records flagged: {totals.get('records_flagged', 0):,}"])
    ws.append([])
    if state.get("summary"):
        ws.append([styled("Executive summary", bold=True)])
        for line in str(state["summary"]).splitlines():
            if line.strip():
                ws.append([line.strip()])
        ws.append([])
    if state.get("catalog_note"):
        ws.append([("WARNING: " if state.get("catalog_stale") else "Note: ") + str(state["catalog_note"])])
        ws.append([])
    finding_rows(state.get("ranked") or [])

    discovered = state.get("discovered_ranked") or []
    if discovered:
        totals_d = state.get("discovered_totals") or {}
        ws.append([])
        ws.append([styled("ON TRIAL - proposed automatically, accepted for trial, NOT counted in the score or in any "
                          "total above", bold=True)])
        finding_rows(discovered)
        ws.append([styled(f"{totals_d.get('probes_run', 0):,} rule(s) on trial · "
                          f"{totals_d.get('records_flagged', 0):,} record(s) flagged · excluded from the score and "
                          f"every headline figure", bold=True)])

    gaps = state.get("empty_scope") or []
    not_running = state.get("not_running") or []
    failed = [r for r in (state.get("results") or []) if not r.ok]
    if gaps or not_running or failed:
        limit = settings.report_reason_max_chars
        ws.append([])
        ws.append([styled("Coverage gaps - these checks did NOT report on their subject", bold=True)])
        for rule_id in gaps:
            ws.append([rule_id, "examined no records at all - this is not a clean result"])
        for result in failed:
            ws.append([result.rule_id, f"failed while running: {result.error}"[:limit]])
        for item in not_running:
            ws.append([item["rule_id"], f"{item['status']}: {item['reason']}"[:limit]])


def _write_rule_sheet(wb, title_for_sheet: str, result, rule) -> int:
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension

    title = getattr(rule, "title", result.rule_id)
    ws = wb.create_sheet(title=title_for_sheet)
    # MUST come before the first append(): in write_only mode widths set afterwards are discarded.
    for i, width in enumerate(content_widths(result.detail_columns, result.detail_rows), 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter] = ColumnDimension(ws, index=letter, width=width)
    ws.freeze_panes = "A2"
    fill = PatternFill("solid", fgColor=SEVERITY_COLOURS.get(getattr(rule, "severity", "medium"), _HEADER_FILL))
    head = []
    for name in result.detail_columns:
        c = WriteOnlyCell(ws, value=header(name))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
        head.append(c)
    ws.append(head)

    body_limit = XLSX_MAX_ROWS - 8
    written = 0
    source = spool_store.read_rows(result.spool_path) if result.spool_path else iter(result.detail_rows)
    for row in source:
        if written >= body_limit:
            break
        ws.append([number_or_text(v) for v in row])
        written += 1
    ws.append([])
    ws.append([f"Rule {result.rule_id}: {title}"])
    ws.append([f"{result.anomaly_count:,} of {result.scope_total:,} records examined were flagged "
               f"({result.anomaly_pct:.2f}%)."])
    if result.spool_truncated:
        ws.append([f"NOTE: this sheet is capped at {cfg().export_max_rows:,} rows (the Excel row limit setting); the "
                   f"rule found more. Rows are ordered worst first, so the most serious are present."])
    if written >= body_limit:
        ws.append([f"NOTE: truncated at {body_limit:,} rows - one worksheet cannot hold more than {XLSX_MAX_ROWS:,}."])
    for concern in result.concerns:
        ws.append([f"NOTE: {concern}"])
    return written


def build(state: dict) -> bytes:
    """The workbook as bytes."""
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    _write_summary_sheet(wb, state)
    results = {r.rule_id: r for r in (state.get("results") or [])}
    rules = state.get("rules") or {}
    for rule_id, title in sheet_names(state).items():
        _write_rule_sheet(wb, title, results[rule_id], rules.get(rule_id))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
