"""The Excel export of one report date (ported from export_service.py; openpyxl instead of xlsxwriter).

Three sheets, as the original: Daily Detail (one row per logical task, Status -> WBS -> Activity -> Well order),
Summary (status -> WBS -> activity, a blank total where the units are mixed) and Data Quality. A pure read, built
in memory.
"""
from __future__ import annotations

import io
import logging
from datetime import date
from decimal import Decimal
from typing import Any, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.features.dailyreport.models import DailyDataset, DailyTask, QuantityStatus

logger = logging.getLogger(__name__)

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

DETAIL_COLUMNS: Sequence[tuple] = (
    ("UOM", 12), ("WBS", 38), ("Activity Code", 18), ("Activity Description", 44), ("Well ID", 10), ("Date", 12),
    ("Task Code", 20), ("Progress", 10), ("Planned Quantity", 17), ("Actual Quantity", 17), ("Quantity Status", 16),
    ("Crew Code", 12), ("Daily Completed", 15), ("PH Name", 20), ("Mapping Status", 20), ("Data Quality Flags", 34),
    # Reference value from the activity master, not authoritative: UOM above is the only unit used.
    ("Activity Master UOM", 16),
)

_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_TITLE = Font(bold=True, size=14)
_NOTE = Font(size=9, italic=True, color="555555")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", start_color="1F3A5F", end_color="1F3A5F")
_GROUP_FILL = PatternFill("solid", start_color="EDF2F8", end_color="EDF2F8")
_GROUP_FONT = Font(bold=True)
_NUMBER = "#,##0.####"
_INTEGER = "0"
_DATE = "yyyy-mm-dd"


class ExportService:
    def __init__(self, ctx) -> None:
        self._ctx = ctx

    def filename(self, report_date: date) -> str:
        return f"{self._ctx.settings.export_file_prefix}_{report_date.isoformat()}.xlsx"

    def build_workbook(self, dataset: DailyDataset) -> bytes:
        workbook = Workbook()
        workbook.remove(workbook.active)
        self._detail_sheet(workbook, dataset)
        self._summary_sheet(workbook, dataset)
        self._data_quality_sheet(workbook, dataset)
        buffer = io.BytesIO()
        workbook.save(buffer)
        content = buffer.getvalue()
        logger.info("excel generated for %s: %d task row(s), %d bytes", dataset.report_date, len(dataset.tasks),
                    len(content))
        return content

    # ------------------------------------------------------------------------------------------------------------
    @staticmethod
    def _heading(sheet, title: str, note: str) -> None:
        sheet.cell(row=1, column=1, value=title).font = _TITLE
        sheet.cell(row=2, column=1, value=note).font = _NOTE

    @staticmethod
    def _header(sheet, row: int, names: Sequence[str]) -> None:
        for index, name in enumerate(names, start=1):
            cell = sheet.cell(row=row, column=index, value=name)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.border = _BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    @staticmethod
    def _widths(sheet, widths: Sequence[int]) -> None:
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    @staticmethod
    def _cell(sheet, row: int, col: int, value: Any, *, group: bool = False) -> None:
        """Missing information stays visibly missing; it is never filled in."""
        cell = sheet.cell(row=row, column=col)
        cell.border = _BORDER
        if group:
            cell.font = _GROUP_FONT
            cell.fill = _GROUP_FILL
        if value is None or value == "":
            cell.alignment = Alignment(vertical="top")
            return
        if isinstance(value, bool):
            cell.value = "Yes" if value else "No"
        elif isinstance(value, date):
            cell.value = value
            cell.number_format = _DATE
        elif isinstance(value, Decimal):
            cell.value = float(value)
            cell.number_format = _NUMBER
        elif isinstance(value, int):
            cell.value = value
            cell.number_format = _INTEGER
        else:
            cell.value = str(value)
            cell.alignment = Alignment(vertical="top")

    # ------------------------------------------------------------------------------------------------------------
    def _detail_sheet(self, workbook: Workbook, dataset: DailyDataset) -> None:
        sheet = workbook.create_sheet("Daily Detail")
        self._heading(sheet, f"Daily Morning Brief - Detail - {dataset.report_date.isoformat()}",
                      "One row per logical daily task for live wells. Quantities are shown in their own unit of "
                      "measure; no conversion between units is applied.")
        header_row = 4
        self._header(sheet, header_row, [name for name, _ in DETAIL_COLUMNS])
        self._widths(sheet, [width for _, width in DETAIL_COLUMNS])

        status_order = {member.value: index for index, member in enumerate(QuantityStatus)}
        ordered = sorted(dataset.tasks, key=lambda t: (status_order[t.quantity_status.value], t.wbs is None,
                                                       t.wbs or "", t.activity_code or "", t.well_id,
                                                       t.task_code or ""))
        row = header_row + 1
        for task in ordered:
            for col, value in enumerate(_task_values(task), start=1):
                self._cell(sheet, row, col, value)
            row += 1

        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
        if row > header_row + 1:
            sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(DETAIL_COLUMNS))}{row - 1}"
        else:
            sheet.cell(row=row, column=1, value="No daily task records for this date.").font = _NOTE

    def _summary_sheet(self, workbook: Workbook, dataset: DailyDataset) -> None:
        sheet = workbook.create_sheet("Summary")
        self._heading(sheet, f"Daily Morning Brief - Summary - {dataset.report_date.isoformat()}",
                      "Grouped Validation Status -> WBS -> Activity. Well count counts distinct wells; task count "
                      "counts tasks. A blank Planned/Actual where the UOM column reads 'mixed' means the group spans "
                      "several units and no conversion between units is defined -- it does not mean zero.")
        header_row = 4
        self._header(sheet, header_row, ["Validation Status", "WBS", "Activity Code", "Activity Description", "UOM",
                                         "Wells", "Tasks", "Planned Quantity", "Actual Quantity"])
        self._widths(sheet, [18, 38, 18, 44, 14, 8, 8, 17, 17])

        row = header_row + 1
        for status_group in self._ctx.grouping.build(dataset):
            status_label = status_group.status.value.replace("_", " ").title()
            for wbs_group in status_group.sorted_wbs_groups():
                for col, value in enumerate([status_label, wbs_group.wbs or "", "All activities", "",
                                             _uom_label(wbs_group), wbs_group.well_count, wbs_group.task_count,
                                             wbs_group.planned_quantity, wbs_group.actual_quantity], start=1):
                    self._cell(sheet, row, col, value, group=col <= 7)
                row += 1
                for activity in wbs_group.sorted_activities():
                    for col, value in enumerate([status_label, wbs_group.wbs, activity.activity_code,
                                                 activity.activity_description, _uom_label(activity),
                                                 activity.well_count, activity.task_count,
                                                 activity.planned_quantity, activity.actual_quantity], start=1):
                        self._cell(sheet, row, col, value)
                    row += 1

        if row == header_row + 1:
            sheet.cell(row=row, column=1, value="No daily task records for this date.").font = _NOTE
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    def _data_quality_sheet(self, workbook: Workbook, dataset: DailyDataset) -> None:
        sheet = workbook.create_sheet("Data Quality")
        self._heading(sheet, f"Data Quality - {dataset.report_date.isoformat()}",
                      "Reported separately from the operational quantities; these conditions do not change any "
                      "planned or actual figure.")
        self._widths(sheet, [34, 14])
        quality = self._ctx.grouping.data_quality(dataset)
        rows: List[tuple] = [
            ("Raw task_daily rows read", quality["raw_row_count"]),
            ("Logical daily tasks after grain resolution", quality["logical_task_count"]),
            ("Rows superseded by grain resolution", quality["superseded_row_count"]),
            ("Tasks with more than one row", quality["multi_row_task_count"]),
            ("Tasks with more than one actual entry", quality["duplicate_actual_task_count"]),
            ("Rows with invalid daily_data JSON", quality["invalid_json_row_count"]),
            ("Rows with unreadable actual_quantity", quality["unparseable_actual_row_count"]),
            ("Tasks carrying at least one flag", quality["affected_task_count"]),
        ]
        header_row = 4
        self._header(sheet, header_row, ["Condition", "Count"])
        row = header_row + 1
        for label, value in rows:
            self._cell(sheet, row, 1, label)
            self._cell(sheet, row, 2, int(value))
            row += 1

        row += 1
        self._header(sheet, row, ["Flag", "Tasks"])
        row += 1
        flag_counts = quality["flag_counts"]
        if flag_counts:
            for flag, count in sorted(flag_counts.items(), key=lambda item: (-item[1], item[0])):
                self._cell(sheet, row, 1, flag)
                self._cell(sheet, row, 2, int(count))
                row += 1
        else:
            self._cell(sheet, row, 1, "No data-quality flags raised.")
            self._cell(sheet, row, 2, 0)


def _task_values(task: DailyTask) -> List[Any]:
    return [task.uom_code, task.wbs, task.activity_code, task.activity_description, task.well_id, task.action_on,
            task.task_code, task.progress, task.planned, task.actual_quantity, task.quantity_status.value,
            task.crew_code, task.daily_completed, task.ph_name, task.mapping_status.value,
            ", ".join(flag.value for flag in task.data_quality_flags), task.activity_uom]


def _uom_label(group) -> str:
    """The group's unit, or the units it spans when no total was reported."""
    if group.quantities_summable:
        return group.uom_code or ""
    return f"mixed ({', '.join(group.uom_codes)})"
