"""COMPILE -> RUN -> DISCOVER on a small fake schema with scripted model replies, through the HTTP endpoints (router
mounted on a bare FastAPI app): the SSE event contract, the result shapes, the catalog round trip, staleness, the
reports (valid .xlsx / .docx), and the error contract."""
from __future__ import annotations

import base64
import io
import json
import zipfile
import xml.dom.minidom

from openpyxl import load_workbook

from tests.features.anomaly import _support as S
from tests.features.anomaly.conftest import PREFIX


def stream(client, path: str, **fields):
    response = client.post(PREFIX + path, json=S.body(**fields))
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/event-stream")
    events = S.parse_sse(response.text)
    terminal = [e for e in events if e[0] in ("result", "error")]
    assert len(terminal) == 1 and events[-1] is terminal[0], [e[0] for e in events]
    for name, data in events:
        assert name in ("progress", "probe", "result", "error", "ping")
        if name == "progress":
            assert set(data) == {"stage", "done", "total", "label"}
        if name == "probe":
            assert set(data) == {"probe"} and data["probe"]["ruleId"]
    return events


def result_of(events) -> dict:
    name, data = events[-1]
    assert name == "result", data
    return data


def compile_all(client, **fields) -> dict:
    return result_of(stream(client, "/compile", rules=S.sample_rules(), catalog=[], **fields))


# ── COMPILE ──

def test_every_compiled_probe_is_streamed_before_the_result(client, fake_llm, fake_db):
    # The Tenant API stores each probe as it arrives, so a compile cut short keeps what it finished.
    events = stream(client, "/compile", rules=S.sample_rules(), catalog=[])
    streamed = {data["probe"]["ruleId"]: data["probe"] for name, data in events if name == "probe"}
    out = result_of(events)
    assert set(streamed) == {p["ruleId"] for p in out["catalog"]}
    assert all(streamed[p["ruleId"]] == p for p in out["catalog"])


def test_a_recompile_streams_nothing_for_reused_probes(client, fake_llm, fake_db):
    first = compile_all(client)
    events = stream(client, "/compile", rules=S.sample_rules(), catalog=first["catalog"])
    assert [d for n, d in events if n == "probe"] == []
    assert sorted(result_of(events)["reused"]) == sorted(first["compiled"])


def test_rules_compile_in_parallel_with_the_same_outcome(client, fake_llm, fake_db):
    parallel = result_of(stream(client, "/compile", rules=S.sample_rules(), catalog=[],
                                settings={"compileConcurrency": 4}))
    assert sorted(parallel["compiled"]) == ["DQ-G01-001", "DQ-G01-002", "DQ-T01", "DQ-T02"]
    assert parallel["failed"] == [] and parallel["stoppedEarly"] == ""
    assert {p["ruleId"]: p["status"] for p in parallel["catalog"]} == {r: "active" for r in parallel["compiled"]}


def test_compile_builds_every_runnable_rule(client, fake_llm, fake_db):
    out = compile_all(client)
    assert set(out) == {"catalog", "compiled", "reused", "failed", "notApplicable", "removed", "duplicates",
                        "llmCalls", "seconds", "stoppedEarly", "errors", "usage"}
    # Two foreign keys -> two family members: one authored and reviewed, the other cloned and smoke-tested.
    assert sorted(out["compiled"]) == ["DQ-G01-001", "DQ-G01-002", "DQ-T01", "DQ-T02"]
    assert out["failed"] == [] and out["notApplicable"] == [] and out["stoppedEarly"] == ""
    by_id = {p["ruleId"]: p for p in out["catalog"]}
    assert set(by_id) == set(out["compiled"])            # draft / rejected rules never compile
    probe = by_id["DQ-T01"]
    assert set(probe) == {"ruleId", "summarySql", "detailSql", "status", "source", "ruleHash", "structureFingerprint",
                          "tableFingerprint", "semanticFingerprint", "compiledAt", "tables", "groundingNote",
                          "verifierNote", "thresholdNote", "error", "llmCalls"}
    assert probe["status"] == "active" and "'DQ-T01' AS rule_id" in probe["summarySql"]
    assert probe["groundingNote"] == "task dates live on task_daily" and probe["verifierNote"]
    assert probe["tableFingerprint"] and probe["structureFingerprint"]
    clones = [p for p in out["catalog"] if "built from the query authored for" in p["groundingNote"]]
    assert len(clones) == 1 and "{{" not in clones[0]["summarySql"] and clones[0]["llmCalls"] == 0
    # The family member is never grounded; the two declared rules are.
    assert fake_llm.count("grounding") == 2 and fake_llm.count("author") == 3 and fake_llm.count("verifier") == 3
    assert out["llmCalls"] == 8
    usage = {u["agent"]: u for u in out["usage"]}
    assert usage["anomaly.grounding"]["model"] == "fast-model" and usage["anomaly.sql_author"]["calls"] == 3
    assert any("DQ-G01-002" in sql and "LEFT JOIN" in sql for sql in fake_db.executed)       # the smoke test


def test_a_second_compile_reuses_what_is_still_valid(client, fake_llm):
    first = compile_all(client)
    calls = len(fake_llm.calls)
    second = result_of(stream(client, "/compile", rules=S.sample_rules(), catalog=first["catalog"]))
    assert sorted(second["reused"]) == sorted(first["compiled"]) and second["compiled"] == []
    assert len(fake_llm.calls) == calls and second["llmCalls"] == 0
    # A changed rule definition rebuilds that rule only; a status change (probation -> active) costs nothing.
    rules = S.sample_rules()
    rules[0]["body"] += "\nNever flag: cancelled tasks."
    rules[1]["status"] = "active"
    third = result_of(stream(client, "/compile", rules=rules, catalog=first["catalog"]))
    assert third["compiled"] == ["DQ-T01"] and "DQ-T02" in third["reused"]


def test_a_rule_that_is_switched_off_loses_its_probe(client):
    first = compile_all(client)
    rules = S.sample_rules()
    rules[1]["status"] = "disabled"
    out = result_of(stream(client, "/compile", rules=rules, catalog=first["catalog"]))
    assert out["removed"] == ["DQ-T02"] and "DQ-T02" not in {p["ruleId"] for p in out["catalog"]}


def test_compile_rejections_are_fed_back_and_retried(client, fake_llm):
    bad = S.TASK_SQL.replace("<ID>", "DQ-T01").replace("ORDER BY severity_value DESC", "")
    fake_llm.script["author"] = [bad, S.TASK_SQL.replace("<ID>", "DQ-T01")]
    out = result_of(stream(client, "/compile", rules=S.sample_rules()[:1], catalog=[]))
    assert out["compiled"] == ["DQ-T01"]
    retry = fake_llm.of("author")[1].user
    assert "IT WAS REJECTED BEFORE IT RAN:" in retry and "DETAIL has no ORDER BY" in retry
    assert "YOUR PREVIOUS ATTEMPT WAS:" in retry and "You have already made 1 attempt(s) at this rule." in retry


def test_the_verifier_fails_closed(client, fake_llm):
    fake_llm.script["verifier"] = "I think it is fine, more or less"
    out = result_of(stream(client, "/compile", rules=S.sample_rules()[:1], catalog=[],
                           settings={"verifyRetries": 0}))
    assert out["failed"] == ["DQ-T01"]
    probe = out["catalog"][0]
    assert probe["status"] == "failed" and "The automated review returned nothing usable" in probe["error"]


def test_the_reviewer_can_retire_a_rule_as_not_applicable(client, fake_llm):
    fake_llm.script["verifier"] = json.dumps({"ok": False, "not_applicable": True, "reason": "no crew data here"})
    out = result_of(stream(client, "/compile", rules=S.sample_rules()[:1], catalog=[]))
    assert out["notApplicable"] == ["DQ-T01"] and out["catalog"][0]["error"] == "no crew data here"


def test_a_probe_that_examines_nothing_is_never_stored_active(client, fake_db):
    fake_db.summaries["DQ-T01"] = (0, 0)
    out = result_of(stream(client, "/compile", rules=S.sample_rules()[:1], catalog=[]))
    assert out["notApplicable"] == ["DQ-T01"] and "examined 0 records" in out["catalog"][0]["error"]


def test_the_circuit_breaker_stops_between_rules(client, fake_llm):
    out = result_of(stream(client, "/compile", rules=S.sample_rules(), catalog=[],
                           settings={"maxCompileCalls": 3}))
    assert "compile budget of 3 LLM calls" in out["stoppedEarly"]
    assert out["compiled"] and len(out["compiled"]) < 4


# ── RUN ──

def test_run_scores_reports_and_keeps_probation_apart(client, fake_llm):
    catalog = compile_all(client)["catalog"]
    out = result_of(stream(client, "/run", rules=S.sample_rules(), catalog=catalog, runId="run-1"))
    for key in ("run_id", "started_at", "database", "seconds", "score", "score_basis", "rule_set_fingerprint",
                "summary", "totals", "by_severity", "by_category", "ranked", "discovered_ranked", "discovered_totals",
                "empty_scope", "not_running", "failed", "catalog_stale", "catalog_note", "error", "staleProbes",
                "reports", "usage"):
        assert key in out, key
    assert out["run_id"] == "run-1" and out["database"] == {"name": "welldb", "server": "test-server"}
    assert out["summary"] == S.SUMMARY and out["error"] == "" and out["catalog_stale"] is False
    # Scored: DQ-T01 (3 of 600) and two clean FK members; DQ-T02 is on probation: reported apart, counted nowhere.
    assert [r["rule_id"] for r in out["ranked"]] == ["DQ-T01"]
    assert out["totals"]["probes_run"] == 3 and out["totals"]["records_flagged"] == 3
    assert [r["rule_id"] for r in out["discovered_ranked"]] == ["DQ-T02"]
    assert out["discovered_totals"]["records_flagged"] == 2
    assert out["score"] < 100 and "critical 40" in out["score_basis"]
    assert {u["agent"] for u in out["usage"]} == {"anomaly.summarizer"}
    summarizer = fake_llm.of("summarizer")[0]
    assert "HEADLINE NUMBERS:" in summarizer.user and "Write the executive summary now." in summarizer.user
    assert "SELECT" not in summarizer.user                       # the summary writer never sees SQL
    # The reports: a valid workbook (a summary sheet + one sheet per rule with findings) and a valid Word package.
    book = load_workbook(io.BytesIO(base64.b64decode(out["reports"]["xlsx"])))
    assert book.sheetnames[0] == "Summary" and any(name.startswith("DQ-T01") for name in book.sheetnames)
    finding = book[next(n for n in book.sheetnames if n.startswith("DQ-T01"))]
    assert finding["A1"].value == "Key" and finding["A2"].value == 101
    docx = zipfile.ZipFile(io.BytesIO(base64.b64decode(out["reports"]["docx"])))
    assert {"[Content_Types].xml", "word/document.xml"} <= set(docx.namelist())
    for name in docx.namelist():
        if name.endswith(".xml") or name.endswith(".rels"):
            xml.dom.minidom.parseString(docx.read(name))          # every part is well-formed XML
    document = docx.read("word/document.xml").decode("utf-8")
    assert "Data Quality Report" in document and "DQ-T01 - Task finished before it started" in document
    assert "Findings from rules on trial" in document
    footer = "".join(docx.read(n).decode("utf-8") for n in docx.namelist() if "footer" in n)
    assert "may contain errors or inaccuracies" in footer


def test_run_without_an_ai_provider_writes_a_plain_summary(client, fake_llm):
    catalog = compile_all(client)["catalog"]
    out = result_of(stream(client, "/run", rules=S.sample_rules(), catalog=catalog, runId="r", llm=None,
                           formats=["xlsx"]))
    assert out["summary"].startswith("1 of 3 checks found problems") and set(out["reports"]) == {"xlsx"}
    assert out["usage"] == []


def test_a_failing_probe_is_reported_not_fatal(client, fake_db):
    catalog = compile_all(client)["catalog"]
    fake_db.failing.add("well.well_master w")
    out = result_of(stream(client, "/run", rules=S.sample_rules(), catalog=catalog, runId="r"))
    assert [f["rule_id"] for f in out["failed"]] == ["DQ-T02"] and "Invalid column name" in out["failed"][0]["error"]


def test_stale_probes_are_never_executed(client, fake_db, fake_llm):
    """Staleness is judged PER PROBE, on the tables it reads (here the schema is pruned per rule)."""
    prune = {"schemaPruneAboveTables": 0}
    catalog = compile_all(client, settings=prune)["catalog"]
    tables = {p["ruleId"]: set(p["tables"]) for p in catalog}
    assert tables["DQ-T02"] == {"well.well_master", "well.task_daily"}           # named + one FK away
    fake_db.columns.append(("well", "task_crew", "remarks", "nvarchar", 200, "YES"))   # the structure moved
    fake_db.executed.clear()
    out = result_of(stream(client, "/run", rules=S.sample_rules(), catalog=catalog, runId="r", settings=prune))
    stale = {s["ruleId"] for s in out["staleProbes"]}
    assert stale == {rule_id for rule_id, read in tables.items() if "well.task_crew" in read}
    assert "DQ-T01" in stale and "DQ-T02" not in stale
    assert all(set(s) == {"ruleId", "title", "reason", "rebuilt"} and not s["rebuilt"] for s in out["staleProbes"])
    assert all(s["reason"] == "the structure of the tables it reads changed" for s in out["staleProbes"])
    assert out["catalog_stale"] is True and "NOT run" in out["catalog_note"]
    assert {n["rule_id"] for n in out["not_running"] if n["status"] == "stale"} == stale
    assert not any("'DQ-T01'" in sql for sql in fake_db.executed)
    assert any("'DQ-T02'" in sql for sql in fake_db.executed)
    assert "catalog" not in out


def test_auto_compile_rebuilds_stale_probes_inside_the_run(client, fake_db):
    catalog = compile_all(client)["catalog"]
    fake_db.columns.append(("well", "task_daily", "remarks", "nvarchar", 200, "YES"))
    out = result_of(stream(client, "/run", rules=S.sample_rules(), catalog=catalog, runId="r",
                           settings={"autoCompile": True}))
    assert out["catalog_stale"] is False and all(s["rebuilt"] for s in out["staleProbes"])
    assert sorted(out["recompiled"]) == sorted(s["ruleId"] for s in out["staleProbes"])
    assert {p["ruleId"] for p in out["catalog"]} == {"DQ-T01", "DQ-T02", "DQ-G01-001", "DQ-G01-002"}
    assert [r["rule_id"] for r in out["ranked"]] == ["DQ-T01"]


def test_auto_compile_respects_its_limit(client, fake_db):
    catalog = compile_all(client)["catalog"]
    fake_db.columns.append(("well", "task_daily", "remarks", "nvarchar", 200, "YES"))
    out = result_of(stream(client, "/run", rules=S.sample_rules(), catalog=catalog, runId="r",
                           settings={"autoCompile": True, "autoCompileMaxRules": 1}))
    assert out["catalog_stale"] is True and "too many to rebuild automatically" in out["catalog_note"]


# ── DISCOVER ──

def test_discover_keeps_only_grounded_new_proposals(client, fake_llm):
    catalog = compile_all(client)["catalog"]
    good = {"title": "Task code names a task that does not exist", "what_is_wrong": "A crew row names a missing task.",
            "why_it_matters": "Crew hours are lost.", "how_to_detect": "Compare crew task references to tasks.",
            "do_not_flag": "", "category": "Reference integrity", "severity": "high", "entity": "task",
            "evidence": "", "observation_id": "obs-001", "confidence": "high"}
    invented = {**good, "title": "Something invented", "observation_id": "OBS-999"}
    covered = {**good, "title": "Task finished before it started", "what_is_wrong": "finished before started",
               "how_to_detect": "compare task start finish", "observation_id": "OBS-001"}
    refused = {**good, "title": "Well drilled twice in one year", "what_is_wrong": "a well is drilled twice",
               "how_to_detect": "count drillings per well per year", "observation_id": "OBS-001"}
    fake_llm.script["scout"] = json.dumps({"proposals": [good, invented, covered, refused]})
    events = stream(client, "/discover", rules=S.sample_rules(), catalog=catalog,
                    proposals={"accepted": [], "rejected": [{"title": "Well drilled twice in one year", "reason": "no"}]})
    out = result_of(events)
    assert set(out) == {"proposals", "dropped", "observations", "error", "seconds", "llmCalls", "usage"}
    assert [p["title"] for p in out["proposals"]] == [good["title"]], out["dropped"]
    kept = out["proposals"][0]
    assert kept["observation_id"] == "OBS-001" and kept["observation_fact"] and kept["evidence"]
    assert len(kept["hash"]) == 12 and kept["database"] == "welldb" and kept["tables"]
    reasons = {d["title"]: d["reason"] for d in out["dropped"]}
    assert reasons == {"Something invented": "evidence not found", "Task finished before it started": "already covered",
                       "Well drilled twice in one year": "previously refused"}
    assert out["observations"] >= 1 and out["llmCalls"] == 1
    scout = fake_llm.of("scout")[0]
    assert "OBS-001" in scout.user and "ALREADY REFUSED" in scout.user and "refused because: no" in scout.user
    assert "Rigs moved on public holidays" in scout.user                       # a rejected RULE is refused too
    assert "Respond with ONLY this JSON" in scout.system and "SELECT" not in scout.user
    assert [d["stage"] for d in S.only(events, "progress")][0] == "context_builder"


def test_proposal_to_rule_formats_the_decision(client):
    proposal = {"title": "Task  names a\nmissing crew", "what_is_wrong": "A task names a crew that does not exist.",
                "why_it_matters": "Crew utilisation is understated.", "how_to_detect": "Compare task crews with crews.",
                "do_not_flag": "", "category": "Reference integrity", "severity": "HIGH", "entity": "task",
                "evidence": "6,054 of 35,749 rows", "observation_id": "OBS-014", "observation_fact": "measured",
                "hash": "abc123", "database": "welldb"}
    response = client.post(PREFIX + "/proposal-to-rule", json=S.body(
        proposal=proposal, status="probation", rules=[{"ruleId": "DQ-S09"}, {"ruleId": "DQ-S10"}]))
    assert response.status_code == 200, response.text
    rule = response.json()["rule"]
    assert rule["ruleId"] == "DQ-S11" and rule["status"] == "probation" and rule["source"] == "discovered"
    assert rule["title"] == "Task names a missing crew" and rule["severity"] == "high" and rule["sqlMode"] == "authored"
    assert rule["body"].splitlines() == ["Wrong: A task names a crew that does not exist.",
                                         "Matters: Crew utilisation is understated.",
                                         "Detect: Compare task crews with crews.",
                                         "Never flag: Nothing has been excluded yet."]
    assert rule["discovery"]["observationId"] == "OBS-014" and rule["discovery"]["discoveredFrom"] == "welldb"
    assert rule["discovery"]["proposalHash"] == "abc123" and rule["ruleHash"]
    first = client.post(PREFIX + "/proposal-to-rule", json=S.body(proposal=proposal, status="rejected",
                                                                 reason="not a problem")).json()["rule"]
    assert first["ruleId"] == "DQ-S01" and first["status"] == "rejected"
    assert first["discovery"]["reason"] == "not a problem"
    taken = client.post(PREFIX + "/proposal-to-rule", json=S.body(proposal=proposal, ruleId="DQ-S10",
                                                                 rules=[{"ruleId": "DQ-S10"}]))
    assert taken.status_code == 400 and "already in use" in taken.json()["detail"]


def test_introspect_summarises_the_database(client, fake_db):
    response = client.post(PREFIX + "/introspect", json=S.body(refresh=False))
    assert response.status_code == 200, response.text
    out = response.json()
    assert out["tables"] == 4 and out["foreignKeys"] == 2 and out["fromCache"] is False
    assert {t["name"] for t in out["tableList"]} == {"well.well_master", "well.task_daily", "well.task_crew",
                                                     "ref.status_lookup"}            # sysdiagrams is excluded
    again = client.post(PREFIX + "/introspect", json=S.body()).json()
    assert again["fromCache"] is True and again["fingerprint"] == out["fingerprint"]
    assert client.post(PREFIX + "/introspect", json=S.body(refresh=True)).json()["fromCache"] is False


# ── The error contract ──

def test_a_missing_setting_is_400_naming_it(client):
    payload = S.body(rules=[], catalog=[])
    payload["settings"].pop("sampleRows")
    response = client.post(PREFIX + "/compile", json=payload)
    assert response.status_code == 400 and "sampleRows" in response.json()["detail"]


def test_a_missing_prompt_is_400_naming_it(client):
    payload = S.body(rules=[], catalog=[])
    payload["prompts"].pop("anomaly.verifier_system")
    response = client.post(PREFIX + "/compile", json=payload)
    assert response.status_code == 400 and "anomaly.verifier_system" in response.json()["detail"]


def test_empty_business_rules_and_examples_are_allowed(client):
    out = compile_all(client, prompts={"anomaly.business_rules": "", "anomaly.few_shots": ""})
    assert len(out["compiled"]) == 4


def test_no_ai_provider_is_503_for_compile_and_discover(client):
    for path in ("/compile", "/discover"):
        response = client.post(PREFIX + path, json=S.body(llm=None, rules=[], catalog=[]))
        assert response.status_code == 503 and "No AI provider" in response.json()["detail"]


def test_an_unconfigured_scope_is_400(client):
    response = client.post(PREFIX + "/run", json=S.body(runId="r", access={"allowedSchemas": []}))
    assert response.status_code == 400 and "not set up yet" in response.json()["detail"]


def test_the_database_down_is_an_error_event(client, fake_db, monkeypatch):
    from app.features.anomaly import dbio

    def down(**_kwargs):
        raise dbio.DatabaseUnavailable("The database could not be reached.")

    monkeypatch.setattr(dbio, "connect", down)
    events = stream(client, "/compile", rules=S.sample_rules(), catalog=[])
    assert events[-1] == ("error", {"message": "The database could not be reached."})


def test_a_bad_rule_is_reported_and_the_rest_compile(client):
    rules = S.sample_rules()[:1] + [{"ruleId": "DQ-BAD", "title": "x", "severity": "urgent"}]
    out = result_of(stream(client, "/compile", rules=rules, catalog=[]))
    assert out["compiled"] == ["DQ-T01"] and any("DQ-BAD: severity is 'urgent'" in e for e in out["errors"])
