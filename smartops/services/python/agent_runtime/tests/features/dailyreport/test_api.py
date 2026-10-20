"""The endpoints end to end (router mounted on a bare FastAPI app; offline): response shapes, refresh, export,
and the error contract (400 bad input / configuration, 404 unknown well, 503 database)."""
from __future__ import annotations

import io
from datetime import date, timedelta

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.features.dailyreport import data
from app.features.dailyreport.api import router
from tests.features.dailyreport.conftest import (
    REPORT_DATE,
    TODAY,
    activity_row,
    body,
    detail_row,
    make_row,
    seed_queries,
)

PREFIX = "/api/agent-runtime/features/dailyreport"
DAY = REPORT_DATE.isoformat()


@pytest.fixture
def client() -> TestClient:
    app = FastAPI()
    app.include_router(router, prefix=PREFIX)
    return TestClient(app)


@pytest.fixture
def day(fake_db):
    fake_db.set(data.DAILY_DETAILS, [
        make_row(task_daily_id=1, well_id=101, planned=10, actual_quantity=10),
        make_row(task_daily_id=2, well_id=101, task_code="FLME1190-101", planned=10, actual_quantity=4,
                 uom_code="m3"),
        make_row(task_daily_id=3, well_id=202, planned=5, actual_quantity="5.0000", wbs=None),
    ])
    fake_db.set(data.DAILY_SUMMARY, [{"raw_row_count": 5, "raw_well_count": 2, "raw_actual_entry_count": 3,
                                      "invalid_json_row_count": 0, "unparseable_actual_row_count": 0,
                                      "logical_task_count": 3, "multi_row_task_count": 1,
                                      "duplicate_actual_task_count": 0, "superseded_row_count": 2}])
    fake_db.set(data.WELL_ACTIVITY, [activity_row(well_id=101), activity_row(well_id=30349)])
    fake_db.set(data.WELL_DETAIL, [detail_row(well_id=30349), detail_row(well_id=101)])
    fake_db.set(data.MILESTONES, [
        {"well_id": 30349, "milestone": "RIG_ON", "deadline_date": TODAY + timedelta(days=2)},
        {"well_id": 101, "milestone": "FLAF", "deadline_date": TODAY - timedelta(days=5)},
        {"well_id": 102, "milestone": "PEGGING", "deadline_date": TODAY - timedelta(days=1)},
    ])
    fake_db.set(data.ACTIVITY_DATES, [{"report_date": REPORT_DATE, "row_count": 5, "well_count": 2}])
    return fake_db


def post(client, path, **fields):
    return client.post(PREFIX + path, json=body(**fields))


class TestDashboard:
    def test_summary(self, client, day):
        out = post(client, "/summary", reportDate=DAY).json()
        assert out["report_date"] == DAY
        assert out["totals"] == {"task_count": 3, "well_count": 2,
                                 "status_counts": {"ON_PLAN": 2, "ABOVE_PLAN": 0, "BELOW_PLAN": 1, "NO_ACTUAL": 0,
                                                   "NOT_VALIDATED": 0}}
        assert [g["status"] for g in out["status_groups"]] == ["ON_PLAN", "BELOW_PLAN"]
        on_plan = out["status_groups"][0]
        assert on_plan["planned_quantity"] == "15" and on_plan["quantities_summable"] is True
        assert [g["wbs"] for g in on_plan["wbs_groups"]] == ["Straightline Welding incl. supports", None]
        assert out["view_mode"] == "detail" and len(out["tasks"]) == 3
        assert out["data_quality"]["superseded_row_count"] == 2
        assert out["tasks"][2]["actual_quantity"] == "5"            # exact decimal text, no float

    def test_view_mode_comes_from_the_settings(self, client, day):
        out = post(client, "/summary", reportDate=DAY, settings={"detailViewTaskThreshold": 2}).json()
        assert out["view_mode"] == "grouped" and out["tasks"] == [] and out["detail_view_task_threshold"] == 2

    def test_refresh_reloads_the_date(self, client, day):
        post(client, "/summary", reportDate=DAY)
        post(client, "/summary", reportDate=DAY)
        assert day.count(data.DAILY_DETAILS) == 1
        post(client, "/summary", reportDate=DAY, refresh=True)
        assert day.count(data.DAILY_DETAILS) == 2

    def test_details(self, client, day):
        out = post(client, "/details", reportDate=DAY).json()
        assert out["task_count"] == 3 and out["tasks"][0]["crew_employees"] == ["Ali", "Omar"]

    def test_group_details(self, client, day):
        out = post(client, "/group-details", reportDate=DAY, status="ON_PLAN").json()
        assert out["filters"] == {"status": "ON_PLAN", "wbs": None, "activity_code": None, "uom": None}
        assert [w["well_id"] for w in out["wells"]] == [101, 202]
        assert out["wells"][0]["planned_quantity"] == "10"
        no_wbs = post(client, "/group-details", reportDate=DAY, wbs="").json()
        assert [t["task_daily_id"] for t in no_wbs["tasks"]] == [3]

    def test_an_unknown_status_is_400(self, client, day):
        response = post(client, "/group-details", reportDate=DAY, status="LATE")
        assert response.status_code == 400 and "Unknown status" in response.json()["detail"]

    def test_well_activity(self, client, day):
        out = post(client, "/well-activity", reportDate=DAY).json()
        assert [w["well_id"] for w in out["wells"]] == [101, 202, 30349]
        assert out["tasks"] == []
        assert out["wells"][0]["today_reported_task_count"] == 2
        one = post(client, "/well-activity", reportDate=DAY, wellId=30349).json()
        assert one["well_count"] == 1 and [t["well_id"] for t in one["tasks"]] == [30349]
        assert "progress" not in one["tasks"][0]

    def test_dates(self, client, day):
        out = post(client, "/dates").json()
        assert out == {"dates": [{"report_date": DAY, "row_count": 5, "well_count": 2}]}
        assert day.calls[-1] == (data.ACTIVITY_DATES, [30])
        assert post(client, "/dates", limit=5000).status_code == 400

    def test_well(self, client, day):
        out = post(client, "/well", reportDate=DAY, wellId=101).json()
        assert out["task_count"] == 2 and out["activity"] is None

    def test_a_quiet_live_well_shows_its_open_work_not_a_404(self, client, day):
        out = post(client, "/well", reportDate=DAY, wellId=30349).json()
        assert out["task_count"] == 0 and out["activity"]["open_task_count"] == 4
        assert [t["task_code"] for t in out["open_tasks"]] == ["FLME1180-31474"]

    def test_a_well_with_no_evidence_is_404(self, client, day):
        response = post(client, "/well", reportDate=DAY, wellId=999)
        assert response.status_code == 404 and "No daily task records for well 999" in response.json()["detail"]

    def test_milestones(self, client, day):
        out = post(client, "/milestones", today=TODAY.isoformat()).json()
        assert out["window_days"] == 7
        assert [(a["well_id"], a["milestone_label"], a["days_remaining"]) for a in out["upcoming"]] == [
            (30349, "Rig-on", 2)]
        assert out["overdue_count"] == 2 and [a["well_id"] for a in out["overdue"]] == [102, 101]
        capped = post(client, "/milestones", today=TODAY.isoformat(), windowDays=1,
                      settings={"milestoneOverdueDisplayLimit": 1}).json()
        assert capped["upcoming"] == [] and capped["overdue_count"] == 2 and len(capped["overdue"]) == 1
        assert post(client, "/milestones", today=TODAY.isoformat(), windowDays=-1).status_code == 400

    def test_export(self, client, day):
        response = post(client, "/export", reportDate=DAY)
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert 'filename="daily_morning_brief_2026-08-01.xlsx"' in response.headers["content-disposition"]
        book = load_workbook(io.BytesIO(response.content))
        assert book.sheetnames == ["Daily Detail", "Summary", "Data Quality"]
        detail = book["Daily Detail"]
        assert detail["A4"].value == "UOM" and detail.max_row == 7          # title, note, blank, header, 3 rows
        assert detail["F5"].value.date() == REPORT_DATE
        assert book["Summary"]["A5"].value == "On Plan"
        quality = book["Data Quality"]
        assert quality["A5"].value == "Raw task_daily rows read" and quality["B5"].value == 5

    def test_export_of_an_empty_day_still_produces_a_workbook(self, client, fake_db):
        response = post(client, "/export", reportDate=DAY, settings={"exportFilePrefix": 'brief"\r\nX'})
        assert response.status_code == 200
        assert 'filename="brief_X_2026-08-01.xlsx"' in response.headers["content-disposition"]
        book = load_workbook(io.BytesIO(response.content))
        assert book["Daily Detail"]["A5"].value == "No daily task records for this date."


class TestAgentEndpoints:
    def test_brief_keeps_the_classic_fields(self, client, day, fake_model):
        fake_model.answers["writer"] = ["Well 30349 reported nothing on the date and has 4 open tasks."]
        response = post(client, "/agent/brief", reportDate=DAY, scope="well", wellId=30349, today=TODAY.isoformat())
        out = response.json()
        for field in ("available", "explanation", "cached", "evidence", "sql_sources", "proof", "trace", "plan",
                      "verification", "verified", "usage", "models", "tool_results"):
            assert field in out
        assert out["verified"] is True and out["report_date"] == DAY
        assert [p["task_code"] for p in out["proof"]] == ["FLME1180-31474"]
        assert [s["file"] for s in out["sql_sources"]] == [data.WELL_ACTIVITY, data.WELL_DETAIL]

    def test_brief_refuses_a_scope_it_does_not_cover(self, client, day):
        assert post(client, "/agent/brief", reportDate=DAY, scope="uom", today=TODAY.isoformat()).status_code == 400
        assert post(client, "/agent/brief", reportDate=DAY, scope="well", today=TODAY.isoformat()).status_code == 400

    def test_brief_without_an_ai_provider_is_available_false(self, client, day):
        out = post(client, "/agent/brief", reportDate=DAY, scope="day", today=TODAY.isoformat(), llm=None).json()
        assert out["available"] is False and out["explanation"] is None and "No AI provider" in out["error"]

    def test_ask(self, client, day, fake_model):
        fake_model.answers["planner"] = ['{"steps": [{"tool": "list_wells", "args": {"filter": '
                                         '"no_report_with_open_work"}}]}']
        fake_model.answers["writer"] = ["Well 30349 reported nothing on the date and has 4 open tasks."]
        out = post(client, "/agent/ask", reportDate=DAY, question="Which wells did not report?",
                   today=TODAY.isoformat()).json()
        assert out["available"] and out["kind"] == "ask" and out["verified"]
        assert [u["agent"] for u in out["usage"]] == ["dailyreport.planner", "dailyreport.writer"]

    def test_the_question_length_comes_from_the_settings(self, client, day):
        assert post(client, "/agent/ask", reportDate=DAY, question="Hi", today=TODAY.isoformat()).status_code == 400
        response = post(client, "/agent/ask", reportDate=DAY, question="Which wells?", today=TODAY.isoformat(),
                        settings={"questionMaxLength": 5})
        assert response.status_code == 400


class TestErrorContract:
    def test_a_missing_prompt_is_400_naming_it(self, client, day, fake_model):
        prompts = body()["prompts"]
        prompts.pop("dailyreport.explain_system")
        payload = body(reportDate=DAY, scope="well", wellId=30349, today=TODAY.isoformat())
        payload["prompts"] = prompts
        response = client.post(PREFIX + "/agent/brief", json=payload)
        assert response.status_code == 400 and "dailyreport.explain_system" in response.json()["detail"]

    def test_a_missing_setting_is_400_naming_it(self, client, day):
        payload = body(reportDate=DAY)
        payload["settings"].pop("detailViewTaskThreshold")
        response = client.post(PREFIX + "/summary", json=payload)
        assert response.status_code == 400 and "detailViewTaskThreshold" in response.json()["detail"]

    def test_a_missing_query_is_400(self, client, day):
        queries = seed_queries()
        queries.pop(data.DAILY_DETAILS)
        response = post(client, "/summary", reportDate=DAY, queries=queries)
        assert response.status_code == 400 and data.DAILY_DETAILS in response.json()["detail"]

    def test_a_query_that_no_longer_returns_a_column_is_400(self, client, day):
        day.columns_override[data.ACTIVITY_DATES] = ["report_date", "row_count"]
        response = post(client, "/dates")
        assert response.status_code == 400 and "well_count" in response.json()["detail"]

    def test_a_table_the_tenant_withdrew_is_400(self, client, day):
        payload = body(reportDate=DAY)
        payload["access"]["excludedTables"] = ["well.well_master"]
        response = client.post(PREFIX + "/summary", json=payload)
        assert response.status_code == 400 and "well.well_master" in response.json()["detail"]

    def test_a_withdrawn_table_applies_to_cached_results_too(self, client, day):
        assert post(client, "/summary", reportDate=DAY).status_code == 200
        payload = body(reportDate=DAY)
        payload["access"]["excludedTables"] = ["well.well_master"]
        assert client.post(PREFIX + "/summary", json=payload).status_code == 400

    def test_a_write_statement_is_refused(self, client, day):
        queries = seed_queries()
        queries[data.ACTIVITY_DATES] = {"sql": "DELETE FROM well.task_daily", "params": [], "columns": []}
        response = post(client, "/dates", queries=queries)
        assert response.status_code == 400 and "read-only" in response.json()["detail"]

    def test_the_database_down_is_503(self, client, day):
        day.failing.add(data.DAILY_DETAILS)
        response = post(client, "/summary", reportDate=DAY)
        assert response.status_code == 503 and "SELECT" not in response.json()["detail"]

    def test_a_report_date_outside_the_bounds_is_400(self, client, day):
        response = post(client, "/summary", reportDate="1999-12-31")
        assert response.status_code == 400 and "outside the allowed range" in response.json()["detail"]

    def test_a_query_asking_for_an_unknown_parameter_is_400(self, client, day):
        queries = seed_queries()
        queries[data.ACTIVITY_DATES] = {**queries[data.ACTIVITY_DATES], "params": ["limit", "tenant_id"]}
        response = post(client, "/dates", queries=queries)
        assert response.status_code == 400 and "tenant_id" in response.json()["detail"]


def test_today_is_required_where_deadlines_are_evaluated(client, day):
    assert post(client, "/milestones").status_code == 422
    assert post(client, "/agent/brief", reportDate=DAY).status_code == 422
    assert isinstance(TODAY, date)
